"""Generate the SYNTHETIC sample resumes in samples/ (PDF + DOCX).

Every person, employer, email and phone number here is invented. Emails use the
reserved example.com domain and phones use the 555 fictional range.

    uv run python scripts/make_sample_resumes.py
"""

from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

OUT = Path(__file__).resolve().parent.parent / "samples"

PEOPLE = [
    {
        "slug": "alex_rivera_data_engineer",
        "name": "Alex Rivera",
        "title": "Senior Data Engineer",
        "contact": "alex.rivera@example.com | (555) 010-2231 | Springfield (fictional)",
        "summary": "Data engineer with 7 years building batch and streaming pipelines on AWS and Azure.",
        "skills": "Python, SQL, Spark, Airflow, dbt, Kafka, Snowflake, PostgreSQL, Docker, Kubernetes, Terraform, CI/CD",
        "jobs": [
            ("Lead Data Engineer, Northwind Analytics (fictional)", "2021 - Present",
             ["Built a Kafka + Spark streaming platform processing 2B events/day on AWS.",
              "Moved 40 Airflow DAGs to dbt models in Snowflake; cut warehouse cost 30%.",
              "Mentoring four engineers; introduced TDD and GitHub Actions."]),
            ("Data Engineer, Contoso Retail (fictional)", "2018 - 2021",
             ["Designed PostgreSQL and Redis backed REST APIs with FastAPI.",
              "Containerised ETL jobs with Docker and deployed them to Kubernetes."]),
        ],
        "education": "B.Sc. Computer Science, Example State University (fictional), 2018",
    },
    {
        "slug": "samira_haddad_ml_engineer",
        "name": "Samira Haddad",
        "title": "Machine Learning Engineer",
        "contact": "samira.haddad@example.com | (555) 014-7780 | Rivertown (fictional)",
        "summary": "ML engineer focused on NLP and LLMs, from research notebooks to production services.",
        "skills": "Python, PyTorch, TensorFlow, scikit-learn, Hugging Face, LangChain, NLP, Computer Vision, Azure, GCP, MLflow",
        "jobs": [
            ("ML Engineer, Fabrikam AI (fictional)", "2022 - Present",
             ["Fine-tuned transformer models with PyTorch and Hugging Face for document classification.",
              "Shipped an LLM retrieval service with LangChain on Azure; p95 latency under 400 ms.",
              "Ran A/B Testing for ranking models with product and data teams."]),
            ("Data Scientist, Tailspin Health (fictional)", "2019 - 2022",
             ["Built scikit-learn and Deep Learning models for computer vision triage.",
              "Presented Data Visualization dashboards in Tableau to stakeholders."]),
        ],
        "education": "M.Sc. Statistics, Example Institute of Technology (fictional), 2019",
    },
    {
        "slug": "morgan_lee_product_designer",
        "name": "Morgan Lee",
        "title": "Product Designer",
        "contact": "morgan.lee@example.com | (555) 019-4402 | Lakeview (fictional)",
        "summary": "Product designer who prototypes in code and works closely with engineering.",
        "skills": "Figma, UX Design, React, TypeScript, JavaScript, Agile, Scrum, Communication, Stakeholder Management",
        "jobs": [
            ("Senior Product Designer, Woodgrove Labs (fictional)", "2020 - Present",
             ["Led UX Design for a B2B dashboard used by 3,000 teams; built a Figma design system.",
              "Prototyped interactions in React and TypeScript with the front-end team.",
              "Facilitated Scrum ceremonies and Stakeholder Management for two squads."]),
        ],
        "education": "B.Des. Interaction Design, Example College of Art (fictional), 2016",
    },
]


def write_docx(p: dict, path: Path) -> None:
    doc = Document()
    doc.add_heading(p["name"], level=0)
    doc.add_paragraph(p["title"])
    doc.add_paragraph(p["contact"])
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(p["summary"])
    doc.add_heading("Skills", level=1)
    doc.add_paragraph(p["skills"])
    doc.add_heading("Experience", level=1)
    for role, dates, bullets in p["jobs"]:
        doc.add_paragraph(f"{role}  ({dates})").runs[0].bold = True
        for b in bullets:
            doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Education", level=1)
    doc.add_paragraph(p["education"])
    doc.core_properties.author = "synthetic sample generator"
    doc.save(path)


def write_pdf(p: dict, path: Path) -> None:
    jobs = "".join(
        f"<p><b>{role}</b> ({dates})</p><ul>{''.join(f'<li>{b}</li>' for b in bullets)}</ul>"
        for role, dates, bullets in p["jobs"]
    )
    html = f"""
    <h1>{p['name']}</h1><p><i>{p['title']}</i></p><p>{p['contact']}</p>
    <h3>Summary</h3><p>{p['summary']}</p>
    <h3>Skills</h3><p>{p['skills']}</p>
    <h3>Experience</h3>{jobs}
    <h3>Education</h3><p>{p['education']}</p>
    <p style="color:#888;font-size:8px">Synthetic sample resume generated for testing. All details are fictional.</p>
    """
    doc = fitz.open()
    page = doc.new_page()
    page.insert_htmlbox(fitz.Rect(50, 50, 545, 800), html, css="* {font-family: sans-serif; font-size: 11px} h1 {font-size: 22px}")
    doc.set_metadata({"author": "synthetic sample generator", "title": f"{p['name']} (synthetic)"})
    doc.save(path)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for p in PEOPLE:
        write_pdf(p, OUT / f"{p['slug']}.pdf")
        write_docx(p, OUT / f"{p['slug']}.docx")
        print("wrote", p["slug"])
